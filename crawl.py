"""Crawl BAN-PT accreditation directories (institusi, prodi) into CSV and Excel.

Both directory pages render a client-side DataTable: every row is delivered by a
single JSON request and pagination happens in the browser. One request per
dataset therefore captures the full table.
"""

import argparse
import csv
import json
import random
import time
import urllib.request
from dataclasses import dataclass
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE = "https://www.banpt.or.id/direktori"
PAGE_SIZE = 100


@dataclass(frozen=True)
class Dataset:
    page_url: str
    data_url: str
    columns: list[str]
    sheet: str


DATASETS = {
    "institusi": Dataset(
        page_url=f"{BASE}/institusi/pencarian_institusi.php",
        data_url=f"{BASE}/model/dir_aipt/get_data_institusi.php",
        columns=[
            "Perguruan Tinggi",
            "Peringkat",
            "No. SK",
            "Tahun SK",
            "Wilayah",
            "Tanggal Kedaluwarsa",
            "Status Kedaluwarsa",
        ],
        sheet="Institusi",
    ),
    "prodi": Dataset(
        page_url=f"{BASE}/prodi/pencarian_prodi.php",
        data_url=f"{BASE}/model/dir_prodi/get_hasil_pencariannew.php",
        columns=[
            "Perguruan Tinggi",
            "Program Studi",
            "Strata",
            "Wilayah",
            "No. SK",
            "Tahun SK",
            "Peringkat",
            "Tanggal Kedaluwarsa",
            "Status Kedaluwarsa",
        ],
        sheet="Prodi",
    ),
}


def fetch_rows(ds: Dataset) -> list[list]:
    time.sleep(random.uniform(2, 5))
    req = urllib.request.Request(
        ds.data_url,
        headers={
            "User-Agent": "Mozilla/5.0 (research crawler; polite single request)",
            "Referer": ds.page_url,
            "X-Requested-With": "XMLHttpRequest",
        },
    )
    with urllib.request.urlopen(req, timeout=120) as resp:
        payload = json.load(resp)
    rows = payload["data"]
    bad = [r for r in rows if len(r) != len(ds.columns)]
    if not rows or bad:
        raise RuntimeError(f"unexpected payload: {len(rows)} rows, {len(bad)} malformed")
    return [[i, *("" if c is None else str(c).strip() for c in r)] for i, r in enumerate(rows, 1)]


def write_csv(path: str, header: list[str], rows: list[list]) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def write_xlsx(path: str, header: list[str], rows: list[list], ds: Dataset, snapshot: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = ds.sheet
    ws.append(header)
    tahun = header.index("Tahun SK")
    for r in rows:
        ws.append([int(c) if j == tahun and c.isdigit() else c for j, c in enumerate(r)])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(header):
        width = max(len(col), *(len(str(r[i])) for r in rows))
        ws.column_dimensions[get_column_letter(i + 1)].width = min(width + 2, 80)

    meta = wb.create_sheet("Info")
    for k, v in (
        ("Sumber", ds.page_url),
        ("Tanggal data", snapshot),
        ("Jumlah baris", len(rows)),
    ):
        meta.append([k, v])
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("dataset", choices=DATASETS)
    ap.add_argument("--date", default=date.today().isoformat(), help="snapshot date YYYY-MM-DD")
    args = ap.parse_args()

    ds = DATASETS[args.dataset]
    header = ["No", *ds.columns]
    rows = fetch_rows(ds)
    base = f"data_akreditasi_{args.dataset}_{args.date}"
    write_csv(f"{base}.csv", header, rows)
    write_xlsx(f"{base}.xlsx", header, rows, ds, args.date)
    pages = -(-len(rows) // PAGE_SIZE)
    print(f"{len(rows)} rows ({pages} pages x {PAGE_SIZE}) -> {base}.csv, {base}.xlsx")


if __name__ == "__main__":
    main()
