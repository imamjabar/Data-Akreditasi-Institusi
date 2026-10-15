"""Crawl the BAN-PT perguruan tinggi list and map dataset names to Kode PT.

1. Query the bianglala autocomplete endpoint for every two-letter term aa..zz.
2. For every Perguruan Tinggi name in the institusi/prodi datasets that is not in
   the list yet, search its cleaned full name, then its distinctive words.
3. Write the distinct list (Kode PT, Nama PT, Original) and a mapping of every
   dataset (Perguruan Tinggi, Wilayah) pair to its Kode PT.

Every request waits a random 2-5 s. Raw responses are cached in .cache/ so an
interrupted run resumes without re-requesting finished terms.
"""

import argparse
import csv
import difflib
import html
import itertools
import json
import os
import random
import re
import string
import time
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from crawl import write_csv

SEARCH_URL = "https://service.banpt.or.id/bianglala/searchPT.php"
TERM_CACHE = ".cache/pt_terms.jsonl"
SEARCH_CACHE = ".cache/pt_search.jsonl"
MANUAL = "kode_pt_manual.csv"
TERMS = ["".join(p) for p in itertools.product(string.ascii_lowercase, repeat=2)]
ITEM = re.compile(r"^(\S*)\s*-\s+(.+)$")
NO_CODE = {"", "-"}
FUZZY_MIN = 0.8

ABBREV = {
    "stikes": "sekolah tinggi ilmu kesehatan",
    "iain": "institut agama islam negeri",
    "uin": "universitas islam negeri",
    "iai": "institut agama islam",
    "stie": "sekolah tinggi ilmu ekonomi",
    "stmik": "sekolah tinggi manajemen informatika dan komputer",
    "stih": "sekolah tinggi ilmu hukum",
    "stia": "sekolah tinggi ilmu administrasi",
    "stikom": "sekolah tinggi ilmu komunikasi",
    "stisip": "sekolah tinggi ilmu sosial dan ilmu politik",
    "stai": "sekolah tinggi agama islam",
    "stit": "sekolah tinggi ilmu tarbiyah",
    "stkip": "sekolah tinggi keguruan dan ilmu pendidikan",
    "stikip": "sekolah tinggi keguruan dan ilmu pendidikan",
    "stis": "sekolah tinggi ilmu syariah",
    "stak": "sekolah tinggi agama kristen",
    "stt": "sekolah tinggi teologi",
    "amik": "akademi manajemen informatika dan komputer",
    "akper": "akademi keperawatan",
    "akbid": "akademi kebidanan",
    "poltekkes": "politeknik kesehatan",
    "kemenkes": "kementerian kesehatan",
    "apikes": "akademi perekam dan informasi kesehatan",
    "atro": "akademi teknik radiodiagnostik dan radioterapi",
    "aktek": "akademi teknik",
}
TYPOS = {
    "univeritas": "universitas",
    "llmu": "ilmu",
    "insitut": "institut",
    "politenik": "politeknik",
    "potiteknik": "politeknik",
    "sumatra": "sumatera",
}
GENERIC = {
    *" ".join(ABBREV.values()).split(),
    "universitas", "institut", "politeknik", "akademi", "sekolah", "tinggi", "ilmu",
    "kesehatan", "agama", "islam", "ekonomi", "teknik", "teknologi", "keperawatan",
    "kebidanan", "farmasi", "manajemen", "informatika", "komputer", "tarbiyah",
    "keguruan", "pendidikan", "syariah", "kristen", "teologi", "kementerian", "dan",
    "kota", "negeri", "indonesia", "nasional", "internasional", "medika", "husada",
    "administrasi", "perawatan", "terapi", "rs",
} - {"radiodiagnostik", "radioterapi"}


# ---------- fetching ----------

def load_cache(path: str) -> dict[str, list[str] | None]:
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as f:
        return {e["term"]: e["results"] for e in map(json.loads, f)}


def fetch_term(term: str) -> list[str] | None:
    """Return results, or None when the server answers 200 with an empty body."""
    url = f"{SEARCH_URL}?{urllib.parse.urlencode({'term': term})}"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (research crawler; polite)"})
    empty = 0
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=60) as resp:
                body = resp.read()
            if not body.strip():
                empty += 1
                if empty == 2:
                    return None
                time.sleep(random.uniform(2, 5))
                continue
            results = json.loads(body)
            if not isinstance(results, list):
                raise ValueError(f"unexpected payload type {type(results).__name__}")
            return results
        except Exception as e:
            if attempt == 2:
                raise RuntimeError(f"term {term!r} failed: {e}") from e
            time.sleep(30 * (attempt + 1))
    raise RuntimeError(f"term {term!r} failed")


class CachedSearch:
    def __init__(self, path: str):
        self.path = path
        self.cache = load_cache(path)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.fetched = 0

    def __call__(self, term: str) -> list[str] | None:
        if term not in self.cache:
            time.sleep(random.uniform(2, 5))
            results = fetch_term(term)
            self.cache[term] = results
            with open(self.path, "a", encoding="utf-8") as f:
                f.write(json.dumps({"term": term, "results": results}, ensure_ascii=False) + "\n")
            self.fetched += 1
            shown = "EMPTY BODY" if results is None else len(results)
            print(f"  fetch #{self.fetched} {term!r}: {shown}", flush=True)
        return self.cache[term]


# ---------- parsing and normalisation ----------

def parse_item(raw: str) -> tuple[str, str, str] | None:
    """(Kode PT, Nama PT, Original) for one source string; codes '-' and '' become ''."""
    original = raw.strip()
    m = ITEM.match(" ".join(original.split()))
    if not original or not m:
        return None
    return ("" if m[1] in NO_CODE else m[1], m[2], original)


def clean(s: str) -> str:
    s = html.unescape(html.unescape(s))
    s = re.sub(r"(Ã‚|Â)(?=[\s\xa0]|$)", "", s).replace("\xa0", " ")
    return " ".join(s.split())


def words(s: str) -> list[str]:
    """Cleaned words with original casing; parentheses dropped, typos fixed."""
    s = re.sub(r"\([^)]*\)", " ", clean(s))
    out = []
    for w in re.sub(r"[^\w`'’&.-]+", " ", s).split():
        fixed = TYPOS.get(w.casefold())
        out.append(fixed.capitalize() if fixed else w)
    return out


def key(s: str) -> str:
    s = re.sub(r"[`'’\"]", "", " ".join(words(s)).casefold())
    out = []
    for t in re.sub(r"[^\w]+", " ", s).split():
        out += ABBREV.get(t, t).split()
    return " ".join(out)


def keys(name: str) -> set[str]:
    """Full key, plus the key without a trailing ', City' part."""
    ks = {key(name)}
    if "," in name:
        ks.add(key(name.rsplit(",", 1)[0]))
    return ks - {""}


def search_terms(name: str) -> list[str]:
    """Full cleaned name, first two distinctive words, longest distinctive word.

    Quotes are removed: the endpoint answers HTTP 500 to terms containing "'".
    """
    ws = [w for w in (re.sub(r"['’\"]", "", w) for w in words(name)) if w]
    full = " ".join(ws).replace(" & ", " ")
    distinct = [w.strip(".,&'`-") for w in ws if w.casefold().strip(".,&'`-") not in GENERIC
                and w.casefold() not in ABBREV and len(w.strip(".,&'`-")) >= 3]
    terms = [full]
    if len(distinct) >= 2:
        terms.append(" ".join(distinct[:2]))
    if distinct:
        terms.append(max(distinct, key=len))
    return list(dict.fromkeys(t for t in terms if len(t) >= 3))


# ---------- index ----------

class Index:
    def __init__(self):
        self.items: set[tuple[str, str, str]] = set()
        self.by_key: dict[str, set[tuple[str, str, str]]] = defaultdict(set)

    def add(self, results: list[str] | None) -> None:
        for raw in results or ():
            item = parse_item(raw)
            if item is None:
                if raw.strip():
                    raise RuntimeError(f"unparseable entry {raw!r}")
                continue
            if item in self.items:
                continue
            self.items.add(item)
            for k in keys(item[1]):
                self.by_key[k].add(item)

    def exact(self, name: str) -> set[tuple[str, str, str]]:
        return set().union(*(self.by_key.get(k, set()) for k in keys(name)))

    def fuzzy(self, name: str) -> tuple[float, set[tuple[str, str, str]]]:
        """Best-scoring entries by key similarity; scores below FUZZY_MIN are skipped."""
        best, hits = FUZZY_MIN, set()
        for a in keys(name):
            sm = difflib.SequenceMatcher(b=a, autojunk=False)
            for k, items in self.by_key.items():
                sm.set_seq1(k)
                if sm.real_quick_ratio() < best or sm.quick_ratio() < best:
                    continue
                score = sm.ratio()
                if score > best + 1e-9 or not hits and score >= best:
                    best, hits = score, set(items)
                elif abs(score - best) <= 1e-9:
                    hits |= items
        return (best, hits) if hits else (0.0, set())


# ---------- outputs ----------

def write_xlsx(path: str, sheet: str, header: list[str], rows: list[list], info: list[tuple]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in ws.iter_cols(min_row=2):
        for cell in col:
            cell.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(header):
        width = max(len(col), *(len(str(r[i])) for r in rows))
        ws.column_dimensions[get_column_letter(i + 1)].width = min(width + 2, 90)
    meta = wb.create_sheet("Info")
    for k, v in info:
        meta.append([k, v])
    wb.save(path)


def collapse_codes(codes: set[str]) -> set[str]:
    codes = {c for c in codes if c}
    return {c for c in codes if not (len(c) < 6 and c.zfill(6) in codes)}


CODE_TYPE = [  # third digit of Kode PT by institution type
    ("akademi komunitas", "6"), ("universitas", "1"), ("institut", "2"),
    ("sekolah tinggi", "3"), ("akademi", "4"), ("politeknik", "5"),
]


def region_match(code: str, wilayah: str) -> bool:
    """Kode PT prefix vs dataset Wilayah: LLDIKTI number, 00 for PTN, 2+even for PTAN, 2+odd for PTAS."""
    prefix = code.zfill(6)[:2]
    if wilayah.isdigit():
        return prefix == wilayah.zfill(2)
    if wilayah == "PTN":
        return prefix == "00"
    if wilayah in ("PTAN", "PTAS"):
        return prefix[0] == "2" and int(prefix[1]) % 2 == (wilayah == "PTAS")
    return False


def choose_code(codes: set[str], wilayah: str, name: str) -> tuple[str, str]:
    """(Kode PT, basis) or ("", "") when the candidates stay ambiguous."""
    if len(codes) == 1:
        return next(iter(codes)), "satu kandidat"
    pool, basis = codes, []
    k = key(name)
    digit = next((d for prefix, d in CODE_TYPE if k.startswith(prefix)), None)
    same_type = {c for c in pool if digit and c.zfill(6)[2] == digit}
    if same_type and same_type != pool:
        pool, basis = same_type, basis + ["jenis PT"]
    same_region = {c for c in pool if region_match(c, wilayah)}
    if same_region and same_region != pool:
        pool, basis = same_region, basis + ["wilayah"]
    if len(pool) == 1:
        return next(iter(pool)), " + ".join(basis)
    return "", ""


def load_refs(paths: list[str]) -> dict[tuple[str, str], set[str]]:
    """(Perguruan Tinggi, Wilayah) -> dataset names it appears in."""
    refs: dict[tuple[str, str], set[str]] = defaultdict(set)
    for path in paths:
        source = "institusi" if "institusi" in path else "prodi"
        with open(path, encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                refs[(row["Perguruan Tinggi"], row["Wilayah"])].add(source)
    return refs


def load_manual() -> dict[str, dict[str, str]]:
    """Reviewed decisions: Perguruan Tinggi -> {Kode PT, Kandidat Kode PT, Catatan}."""
    if not os.path.exists(MANUAL):
        return {}
    with open(MANUAL, encoding="utf-8") as f:
        return {r["Perguruan Tinggi"]: r for r in csv.DictReader(f)}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", default=date.today().isoformat(), help="snapshot date YYYY-MM-DD")
    args = ap.parse_args()

    terms = CachedSearch(TERM_CACHE)
    print(f"bigram terms: {sum(t in terms.cache for t in TERMS)} cached", flush=True)
    for t in TERMS:
        terms(t)
    empty_terms = [t for t in TERMS if terms.cache[t] is None]

    index = Index()
    for results in terms.cache.values():
        index.add(results)

    ref_paths = [f"data_akreditasi_{d}_{args.date}.csv" for d in ("institusi", "prodi")]
    refs = load_refs(ref_paths)
    manual = load_manual()
    names = sorted({n for n, _ in refs})

    search = CachedSearch(SEARCH_CACHE)
    for results in search.cache.values():
        index.add(results)
    missing = [n for n in names if n not in manual and not index.exact(n)]
    print(f"{len(names)} dataset names, {len(missing)} not in list -> searching", flush=True)
    for n, name in enumerate(missing, 1):
        print(f"[{n}/{len(missing)}] {name}", flush=True)
        for term in search_terms(name):
            index.add(search(term))
            if index.exact(name):
                break

    pt_rows = sorted(index.items, key=lambda r: (r[0] == "", r[0], r[1].casefold(), r[2]))
    pt_rows = [list(r) for r in pt_rows]
    base = f"data_perguruan_tinggi_{args.date}"
    pt_header = ["Kode PT", "Nama PT", "Original"]
    write_csv(f"{base}.csv", pt_header, pt_rows)
    write_xlsx(f"{base}.xlsx", "Perguruan Tinggi", pt_header, pt_rows, [
        ("Sumber", f"{SEARCH_URL}?term=aa .. ?term=zz, plus name searches"),
        ("Tanggal data", args.date),
        ("Jumlah baris", len(pt_rows)),
    ])

    by_code = defaultdict(set)
    for kode, nama, _ in index.items:
        by_code[kode].add(nama)
    map_rows = []
    for (name, wilayah), sources in sorted(refs.items(), key=lambda kv: (kv[0][0].casefold(), kv[0][1])):
        note = ""
        if name in manual:
            m = manual[name]
            codes = {c.strip() for c in m["Kandidat Kode PT"].split(";") if c.strip()}
            chosen, basis = (m["Kode PT"], "manual") if m["Kode PT"] else ("", "")
            method, note = "manual", m["Catatan"]
            matched = set().union(*(by_code.get(c, set()) for c in codes))
        else:
            hits, method = index.exact(name), "exact"
            if not hits:
                score, hits = index.fuzzy(name)
                method = f"fuzzy {score:.2f}" if hits else ""
            codes = collapse_codes({h[0] for h in hits})
            chosen, basis = choose_code(codes, wilayah, name)
            matched = {h[1] for h in hits if not chosen or h[0].zfill(6) == chosen.zfill(6)}
            if not hits:
                note = "tidak ditemukan di daftar BAN-PT"
            elif not codes:
                note = "di daftar BAN-PT tanpa kode"
        map_rows.append([
            name, wilayah, chosen, "; ".join(sorted(codes)), "; ".join(sorted(matched)),
            method, basis, note, "; ".join(sorted(sources)),
        ])
    map_base = f"data_kode_pt_{args.date}"
    map_header = ["Perguruan Tinggi", "Wilayah", "Kode PT", "Kandidat Kode PT", "Nama PT (BAN-PT)",
                  "Metode", "Dasar Kode", "Catatan", "Sumber"]
    write_csv(f"{map_base}.csv", map_header, map_rows)
    write_xlsx(f"{map_base}.xlsx", "Kode PT", map_header, map_rows, [
        ("Sumber", "Perguruan Tinggi dari data_akreditasi_institusi dan data_akreditasi_prodi"),
        ("Tanggal data", args.date),
        ("Jumlah baris", len(map_rows)),
    ])

    methods = defaultdict(int)
    for r in map_rows:
        methods[r[5].split()[0] if r[5] else "not found"] += 1
    print(f"empty-body bigrams: {len(empty_terms)}; new requests: {terms.fetched + search.fetched}")
    print(f"{base}: {len(pt_rows)} rows")
    print(f"{map_base}: {len(map_rows)} rows, {dict(methods)}, "
          f"with Kode PT {sum(bool(r[2]) for r in map_rows)}, "
          f"ambiguous {sum(not r[2] and bool(r[3]) for r in map_rows)}")


if __name__ == "__main__":
    main()
